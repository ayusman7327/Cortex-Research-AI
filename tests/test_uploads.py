"""Generated upload fixtures exercise parsers without live APIs or private data."""
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import Mock

from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageDraw
from pypdf import PdfReader, PdfWriter
import pytest
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

from rag import document_processor as processor


def upload(name, data):
    stream = BytesIO(data)
    stream.name = name
    return stream


def make_pdf(scanned=False):
    stream = BytesIO()
    document = canvas.Canvas(stream)
    if scanned:
        picture = Image.open(BytesIO(make_image()))
        document.drawImage(ImageReader(picture), 40, 600, width=400, height=100)
        picture.close()
    else:
        document.drawString(40, 750, "The research experiment achieved 92 percent accuracy on a public dataset.")
        document.showPage()
        document.drawString(40, 750, "The limitations include a small dataset and no independent replication.")
    document.save()
    return stream.getvalue()


def make_image(format="PNG", size=(800, 200)):
    stream = BytesIO()
    with Image.new("RGB", size, "white") as picture:
        ImageDraw.Draw(picture).text((20, 20), "Accuracy: 92 percent", fill="black")
        picture.save(stream, format=format)
    return stream.getvalue()


@pytest.fixture(autouse=True)
def no_live_api(monkeypatch):
    """Fail immediately if a parser unexpectedly attempts a real API request."""
    def refuse_client():
        raise AssertionError("A live Gemini client must not be used in upload tests")
    monkeypatch.setattr(processor, "get_client", refuse_client)


@pytest.fixture
def transcription_client(monkeypatch):
    client = Mock()
    client.models.generate_content.return_value = SimpleNamespace(
        text="Accuracy: 92 percent", candidates=[SimpleNamespace(finish_reason="STOP")]
    )
    monkeypatch.setattr(processor, "get_client", lambda: client)
    return client


def test_native_pdf_preserves_pages_and_source():
    pages, warnings = processor.extract_document(upload("study.PDF", make_pdf()))
    assert [page["page_number"] for page in pages] == [1, 2]
    assert [page["location"] for page in pages] == ["Page 1", "Page 2"]
    assert all(page["source"] == "study.PDF" for page in pages)
    assert "92 percent" in pages[0]["text"]
    assert "replication" in pages[1]["text"]
    assert warnings == []


def test_docx_preserves_paragraph_table_order():
    document = Document()
    document.add_paragraph("Method: compare two models.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Model", "Accuracy"
    table.cell(1, 0).text, table.cell(1, 1).text = "Cortex", "92%"
    document.add_paragraph("Conclusion: reproduce the experiment.")
    stream = BytesIO()
    document.save(stream)
    pages, _ = processor.extract_document(upload("study.docx", stream.getvalue()))
    text = "\n".join(page["text"] for page in pages)
    assert text.index("Method:") < text.index("Model | Accuracy") < text.index("Cortex | 92%") < text.index("Conclusion:")
    assert pages[0]["location"] == "Section 1"


@pytest.mark.parametrize("name,encoding,text", [
    ("notes.txt", "utf-8", "Accuracy is 92 percent. Café research."),
    ("notes.txt", "utf-8-sig", "A UTF-8 document with a byte order mark."),
    ("notes.txt", "utf-16", "Research notes: café and 東京."),
    ("notes.md", "utf-8", "# Research\n\nThe result is **92 percent**."),
])
def test_text_formats_and_encodings(name, encoding, text):
    pages, warnings = processor.extract_document(upload(name, text.encode(encoding)))
    assert pages[0]["text"] == text
    assert pages[0]["source"] == name
    assert warnings == []


@pytest.mark.parametrize("name,data", [
    ("results.csv", b'Model,Accuracy,Notes\nCortex,92,"public, replicated"\n'),
    ("results.tsv", b'Model\tAccuracy\tNotes\nCortex\t92\tpublic, replicated\n'),
])
def test_delimited_tables(name, data):
    pages, _ = processor.extract_document(upload(name, data))
    assert "Row 1: Model | Accuracy | Notes" in pages[0]["text"]
    assert "Row 2: Cortex | 92 | public, replicated" in pages[0]["text"]
    assert pages[0]["location"] == "Table, rows 1–2"


def test_json_preserves_nested_values_and_unicode():
    data = '{"dataset": "東京", "results": [{"accuracy": 92}], "replicated": true}'
    pages, _ = processor.extract_document(upload("results.json", data.encode("utf-8")))
    assert '"dataset": "東京"' in pages[0]["text"]
    assert '"accuracy": 92' in pages[0]["text"]
    assert '"replicated": true' in pages[0]["text"]


def test_xlsx_multiple_sheets_and_row_locations():
    book = Workbook()
    results = book.active
    results.title = "Results"
    results.append(["Model", "Accuracy"])
    for number in range(1, 42):
        results.append([f"Model {number}", number])
    methods = book.create_sheet("Methods")
    methods.append(["Method", "Cross validation"])
    stream = BytesIO()
    book.save(stream)
    book.close()
    pages, warnings = processor.extract_document(upload("results.xlsx", stream.getvalue()))
    assert [page["location"] for page in pages] == [
        "Sheet Results, rows 1–40", "Sheet Results, rows 41–42", "Sheet Methods, rows 1–1"
    ]
    assert "Row 42: Model 41 | 41" in pages[1]["text"]
    assert "Model | Accuracy" in pages[1]["text"]
    assert "Cross validation" in pages[2]["text"]
    assert any("formulas use saved results" in warning for warning in warnings)


@pytest.mark.parametrize("name,format", [
    ("figure.png", "PNG"), ("figure.jpg", "JPEG"),
    ("figure.jpeg", "JPEG"), ("figure.webp", "WEBP"),
])
def test_images_use_mocked_transcription(name, format, transcription_client):
    pages, warnings = processor.extract_document(upload(name, make_image(format)))
    assert pages[0]["text"] == "Accuracy: 92 percent"
    assert pages[0]["source"] == name
    assert pages[0]["location"] == "Image section 1"
    assert any("AI transcription" in warning for warning in warnings)
    arguments = transcription_client.models.generate_content.call_args.kwargs
    image_part = arguments["contents"][0]
    assert image_part.inline_data.mime_type == "image/png"
    with Image.open(BytesIO(image_part.inline_data.data)) as normalized:
        assert normalized.mode == "RGB"
        assert max(normalized.size) <= 2400
    transcription_client.close.assert_called_once()


def test_scanned_pdf_transcribes_and_keeps_page_reference(transcription_client):
    pages, warnings = processor.extract_document(upload("scanned.pdf", make_pdf(scanned=True)))
    assert pages == [{"text": "Accuracy: 92 percent", "source": "scanned.pdf",
                      "page_number": 1, "location": "Page 1"}]
    assert any("AI read 1 scanned page" in warning for warning in warnings)
    transcription_client.models.generate_content.assert_called_once()
    transcription_client.close.assert_called_once()


@pytest.mark.parametrize("name,data", [
    ("figure.png", make_image()), ("scan.pdf", make_pdf(scanned=True))
])
def test_scans_and_images_require_ai_reading(name, data):
    with pytest.raises(ValueError, match="[Aa][Ii] reading"):
        processor.extract_document(upload(name, data), use_ocr=False)


@pytest.mark.parametrize("name,data", [
    ("broken.pdf", b"not a PDF"), ("broken.docx", b"not a Word file"),
    ("broken.xlsx", b"not an Excel file"), ("broken.png", b"not an image"),
    ("broken.json", b'{"accuracy":'), ("empty.txt", b""),
    ("blank.txt", b" \n\t "), ("program.exe", b"unsupported program"),
    ("binary.txt", b"Data\x00binary"), ("invalid.txt", b"\xff\xff\xff"),
])
def test_unreadable_empty_and_unsupported_files(name, data):
    with pytest.raises(ValueError):
        processor.extract_document(upload(name, data))


def test_encrypted_pdf_has_actionable_error():
    writer = PdfWriter()
    reader = PdfReader(BytesIO(make_pdf()))
    for page in reader.pages:
        writer.add_page(page)
    writer.encrypt("test-only-password")
    stream = BytesIO()
    writer.write(stream)
    with pytest.raises(ValueError, match="Password-protected.*unlocked"):
        processor.extract_document(upload("locked.pdf", stream.getvalue()))


def test_blank_pdf_is_rejected():
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    stream = BytesIO()
    writer.write(stream)
    with pytest.raises(ValueError, match="No readable content"):
        processor.extract_document(upload("blank.pdf", stream.getvalue()))


def test_mixed_batch_retains_valid_documents():
    result = processor.process_uploads([
        upload("good.txt", b"A usable research result."),
        upload("bad.pdf", b"damaged PDF"),
        upload("second.md", b"# Another usable research result"),
    ])
    assert [page["source"] for page in result["pages"]] == ["good.txt", "second.md"]
    assert len(result["errors"]) == 1
    assert result["errors"][0].startswith("bad.pdf:")


def test_duplicate_contents_are_skipped():
    result = processor.process_uploads([
        upload("one.txt", b"Same useful text."), upload("two.txt", b"Same useful text.")
    ])
    assert len(result["pages"]) == 1
    assert result["errors"] == []
    assert any("duplicate contents skipped" in warning for warning in result["warnings"])


def test_invalid_file_does_not_suppress_valid_file_with_same_bytes():
    contents = b"A useful text source with an incorrectly named copy."
    result = processor.process_uploads([
        upload("wrong-extension.pdf", contents), upload("correct.txt", contents)
    ])
    assert [page["source"] for page in result["pages"]] == ["correct.txt"]
    assert len(result["errors"]) == 1


@pytest.mark.parametrize("name,data", [
    ("blank.csv", b",,\n,,\n"), ("blank.tsv", b"\t\t\n\t\t\n")
])
def test_blank_tables_are_rejected(name, data):
    with pytest.raises(ValueError, match="No readable content"):
        processor.extract_document(upload(name, data))


def test_blank_workbook_is_rejected():
    book = Workbook()
    stream = BytesIO()
    book.save(stream)
    book.close()
    with pytest.raises(ValueError, match="No readable content"):
        processor.extract_document(upload("blank.xlsx", stream.getvalue()))


def test_duplicate_filenames_receive_distinct_source_references():
    result = processor.process_uploads([
        upload("study.txt", b"First study result."), upload("study.txt", b"Second study result.")
    ])
    assert [page["source"] for page in result["pages"]] == ["study.txt", "study (2).txt"]
    assert result["errors"] == []


def test_source_names_remain_unique_when_uploaded_names_include_suffixes():
    result = processor.process_uploads([
        upload("study.txt", b"First study result."),
        upload("study.txt", b"Second study result."),
        upload("study (2).txt", b"Third study result."),
    ])
    sources = [page["source"] for page in result["pages"]]
    assert len(sources) == len(set(sources)) == 3


def test_paths_are_removed_from_source_names():
    pages, _ = processor.extract_document(upload(r"C:\private\research\study.txt", b"Public test content."))
    assert pages[0]["source"] == "study.txt"


@pytest.mark.parametrize("count", [0, 11])
def test_batch_file_count_limit(count):
    with pytest.raises(ValueError, match="between 1 and 10"):
        processor.process_uploads([upload(f"{number}.txt", b"text") for number in range(count)])


def test_file_and_batch_size_limits(monkeypatch):
    monkeypatch.setattr(processor, "MAX_FILE_BYTES", 20)
    with pytest.raises(ValueError, match="20 MB"):
        processor.extract_document(upload("large.txt", b"x" * 21))
    monkeypatch.setattr(processor, "MAX_TOTAL_BYTES", 30)
    with pytest.raises(ValueError, match="50 MB"):
        processor.process_uploads([upload("a.txt", b"a" * 16), upload("b.txt", b"b" * 16)])


@pytest.mark.parametrize("text,finish_reason,expected", [
    ("NO_READABLE_CONTENT", "STOP", "No readable content"),
    (None, "STOP", "no transcription"),
    ("partial transcription", "MAX_TOKENS", "exceeded its limit"),
])
def test_image_transcription_failure_closes_client(transcription_client, text, finish_reason, expected):
    transcription_client.models.generate_content.return_value = SimpleNamespace(
        text=text, candidates=[SimpleNamespace(finish_reason=finish_reason)]
    )
    with pytest.raises(ValueError, match=expected):
        processor.extract_document(upload("figure.png", make_image()))
    transcription_client.close.assert_called_once()


def test_image_api_error_keeps_successful_batch_and_hides_internal_details(transcription_client):
    error = RuntimeError("private request payload")
    error.code = 429
    transcription_client.models.generate_content.side_effect = error
    result = processor.process_uploads([
        upload("good.txt", b"A valid source."), upload("figure.png", make_image())
    ])
    assert len(result["pages"]) == 1
    assert "quota" in result["errors"][0].lower()
    assert "private request payload" not in result["errors"][0]
    transcription_client.close.assert_called_once()
